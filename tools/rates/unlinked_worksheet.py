#!/usr/bin/env python3
"""A worksheet for the LCA curators: the Books of Rates entries HECTOR cannot link to a commodity.

    .venv/bin/python -m tools.rates.unlinked_worksheet --out ~/Downloads/books-of-rates-unlinked.xlsx

Uses link_rates' own matching (link, form_index, LINK_OVERRIDES) and preconditions, so it lists
exactly what the linker leaves out for want of a commodity: rows with a price and a unit whose
goods match NO glossary spelling; or only a spelling two or more concepts share (ambiguous); or
only a word naming the material, purpose or container ("Saddels of stele"), which is never taken
for the goods; or that were read by hand and found to have no concept yet (LINK_OVERRIDES None).
Rows lacking a price or a unit are not a curator's question and are left out.

One row per distinct goods phrase (an entry repeated across the books needs one decision), with
every book and line it occurs in, and suggestions: for an ambiguous phrase, the concepts that
share the spelling; otherwise up to three concepts whose spellings are closest to its leading
words. The curator's answer goes in "decision": a suggestion number, a concept key, "new" (with a
name in the note), or "not goods".
"""
import argparse
import collections
import difflib
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from tools.rates import link_rates as L


def suggestions(text: str, idx: dict, spellings: list[str], entries: dict, n: int = 3):
    toks = L.norm(text).split()
    seen, out = set(), []
    for k in (3, 2, 1):
        probe = " ".join(toks[:k])
        if len(probe) < 3:
            continue
        for sp in difflib.get_close_matches(probe, spellings, n=8, cutoff=0.72):
            for key in sorted(idx[sp]):
                if key not in seen:
                    seen.add(key)
                    out.append((key, sp))
        if len(out) >= n:
            break
    return out[:n]


def label(entries, key):
    d = (entries[key].get("d") or "").split("|")[0].strip()
    return f"{key}: {d[:90]}" if d else key


def build(out: Path):
    lca = L.LCA
    entries = json.loads((lca / "docs/data/glossary_data.json").read_text(encoding="utf-8"))["entries"]
    idx = L.form_index(entries)
    max_len = max(len(k.split()) for k in idx)
    spellings = sorted(idx)
    units = {r["key"] for r in L.load_tsv(L.BUILD / "ledger/units.tsv") if r["status"] == "active"}
    join = {r["rates_unit"]: r["key"] for r in L.load_tsv(L.BUILD / "units/rates_join.tsv")}
    groups = collections.OrderedDict()
    counts = collections.Counter()
    for line in (L.BUILD / "rates/rates.jsonl").open(encoding="utf-8"):
        r = json.loads(line)
        if r["book"] not in L.LINKED_BOOKS or not r.get("pence") or join.get(r.get("unit") or "") not in units:
            continue
        text = r.get("commodity_text") or r.get("commodity_raw") or ""
        ov = (r["book"], r["direction"], str(r["source_line"]))
        if ov in L.LINK_OVERRIDES:
            if L.LINK_OVERRIDES[ov]:
                continue                                     # linked, by hand
            key, why = None, "no concept yet: read by hand"
        else:
            key, _sp, _q, why = L.link(text, idx, max_len)
        if key:
            continue
        kind = why.split(":")[0]
        counts[kind] += 1
        g = groups.setdefault(L.norm(text), {"text": text, "kind": kind, "why": why, "where": [], "rate": []})
        g["where"].append(f"{r['book']} l.{r['source_line']}")
        g["rate"].append(f"{r['book']}: {r.get('commodity_raw')} {r.get('rate_raw') or ''}".strip())

    wb = Workbook()
    ws = wb.active
    ws.title = "unlinked"
    head = ["#", "goods as written", "problem", "suggestion 1", "suggestion 2", "suggestion 3",
            "decision", "note", "occurs in", "entries in full"]
    ws.append(head)
    for i, g in enumerate(sorted(groups.values(), key=lambda g: (-len(g["where"]), L.norm(g["text"]))), 1):
        if g["kind"] == "ambiguous":
            words = [w.strip() for w in g["why"].split(":", 1)[1].split(",")]
            cands = [(k, w) for w in words for k in sorted(idx.get(w, ()))][:3]
            problem = f"'{', '.join(words)}' is a spelling of more than one concept"
        elif g["kind"] == "only a material, purpose or container":
            words = [w.strip() for w in g["why"].split(":", 1)[1].split(",")]
            cands = suggestions(g["text"], idx, spellings, entries)
            problem = (f"only '{', '.join(words)}' matched, which names what the goods are made of, for or "
                       "packed in, not the goods themselves")
        elif g["kind"] == "no concept yet":
            cands = suggestions(g["text"], idx, spellings, entries)
            problem = "read by hand: the automatic link was wrong, and no concept was found for the goods by name (a suggestion may still fit)"
        else:
            cands = suggestions(g["text"], idx, spellings, entries)
            problem = "no glossary spelling"
        sugg = [f"{label(entries, k)}  [via '{sp}']" for k, sp in cands] + [""] * 3
        ws.append([i, g["text"], problem, *sugg[:3], "", "", "; ".join(g["where"]), "\n".join(g["rate"])])
    widths = [5, 30, 26, 38, 38, 38, 14, 28, 22, 60]
    for col, w in zip("ABCDEFGHIJ", widths):
        ws.column_dimensions[col].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8E6DF")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=7):
        row[0].fill = PatternFill("solid", fgColor="FFF4D6")
    dv = DataValidation(type="list", formula1='"1,2,3,new,not goods"', allow_blank=True, showErrorMessage=False)
    dv.add(f"G2:G{ws.max_row}")
    ws.add_data_validation(dv)
    ws.freeze_panes = "B2"

    notes = wb.create_sheet("how to fill this in")
    for line in [
        "Books of Rates entries HECTOR cannot yet link to a commodity in the London Customs Accounts glossary.",
        "Each row is one way of writing the goods; 'occurs in' lists every book and line that uses it.",
        "",
        "In the yellow 'decision' column, please write:",
        "  1, 2 or 3   the suggestion that is right (the spelling will be added to that concept);",
        "  a concept key, if none of the suggestions is right but another concept is;",
        "  new         if the glossary has no concept for it (give a name in 'note');",
        "  not goods   if the entry is not a commodity at all.",
        "",
        "'problem' says why it did not link:",
        "  no glossary spelling: no spelling of the goods is in the glossary;",
        "  a spelling of more than one concept (e.g. 'grayne' is both a grain and grain the dyestuff): say which is meant here;",
        "  only a material, purpose or container matched ('Saddels of stele': steel is what the saddles are made of):",
        "     the goods themselves (saddles) have no glossary spelling;",
        "  read by hand: the automatic link was wrong, and no concept was found for the goods by name; a suggestion may",
        "     still fit ('Playne yrones for carpenters': planing iron?), otherwise 'new'.",
        "",
        f"{counts['no glossary spelling']} entries with no glossary spelling, {counts['ambiguous']} ambiguous, "
        f"{counts['only a material, purpose or container']} with only a material, purpose or container, and "
        f"{counts['no concept yet']} read by hand, in {len(groups)} distinct phrases (Books of Rates 1507, 1545, 1558).",
        "Suggestions are by closeness of spelling only and will often be wrong: please check every one.",
    ]:
        notes.append([line])
    notes.column_dimensions["A"].width = 110
    wb.move_sheet("how to fill this in", offset=-1)
    wb.active = 1
    wb.save(out)
    return counts, len(groups)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--out", type=Path, required=True)
    a = ap.parse_args(argv)
    counts, n = build(a.out.expanduser())
    print(f"{dict(counts)} in {n} distinct phrases -> {a.out}")


if __name__ == "__main__":
    main()
